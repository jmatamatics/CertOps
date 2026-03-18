"""Content ingestion utilities for Build Your Own certifications.

Fetches URLs, parses PDF/DOCX uploads, chunks text, and optionally embeds
into Qdrant for program-scoped RAG retrieval.
"""

import io
import os
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from docx import Document
from langchain_openai import OpenAIEmbeddings
from langchain_qdrant import QdrantVectorStore
from langchain_core.documents import Document as LCDocument
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams

load_dotenv(Path(__file__).resolve().parent.parent / ".env")

QDRANT_URL = os.getenv("QDRANT_URL")
QDRANT_API_KEY = os.getenv("QDRANT_API_KEY")
CUSTOM_COLLECTION = "certops_custom_docs"

_embeddings = OpenAIEmbeddings(model="text-embedding-3-small")


def fetch_url(url: str, timeout: float = 15.0) -> str:
    """Fetch a URL and extract readable text from the HTML."""
    resp = httpx.get(url, timeout=timeout, follow_redirects=True)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")

    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()

    text = soup.get_text(separator="\n", strip=True)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def chunk_text(text: str, max_chars: int = 2000) -> list[str]:
    """Split text into chunks, breaking on paragraph boundaries."""
    paragraphs = text.split("\n\n")
    chunks: list[str] = []
    current = ""

    for para in paragraphs:
        if len(current) + len(para) + 2 > max_chars and current:
            chunks.append(current.strip())
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para

    if current.strip():
        chunks.append(current.strip())

    return chunks


def process_content(
    urls: list[str],
    file_contents: list[tuple[str, bytes]],
) -> list[str]:
    """Process URLs and uploaded files into text chunks for the pipeline.

    Args:
        urls: List of web URLs to fetch.
        file_contents: List of (filename, file_bytes) tuples.

    Returns:
        List of text chunks ready for the LangGraph pipeline.
    """
    all_text: list[str] = []

    for url in urls:
        try:
            text = fetch_url(url)
            if text:
                all_text.append(text)
        except Exception as e:
            all_text.append(f"[Error fetching {url}: {e}]")

    for filename, content in file_contents:
        lower = filename.lower()
        try:
            if lower.endswith(".pdf"):
                text = parse_pdf(content)
            elif lower.endswith(".docx"):
                text = parse_docx(content)
            else:
                text = content.decode("utf-8", errors="replace")

            if text:
                all_text.append(text)
        except Exception as e:
            all_text.append(f"[Error parsing {filename}: {e}]")

    chunks: list[str] = []
    for text in all_text:
        chunks.extend(chunk_text(text))

    return chunks


def _ensure_custom_collection() -> None:
    """Create the custom docs collection in Qdrant if it doesn't exist."""
    client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)
    existing = [c.name for c in client.get_collections().collections]
    if CUSTOM_COLLECTION not in existing:
        client.create_collection(
            collection_name=CUSTOM_COLLECTION,
            vectors_config=VectorParams(size=1536, distance=Distance.COSINE),
        )


def embed_program_docs(program_id: str, chunks: list[str], sources: list[str] | None = None) -> int:
    """Embed text chunks into Qdrant with program_id metadata for scoped retrieval.

    Returns the number of chunks embedded.
    """
    if not chunks:
        return 0

    _ensure_custom_collection()

    documents = [
        LCDocument(
            page_content=chunk,
            metadata={
                "program_id": program_id,
                "source": (sources[min(i, len(sources) - 1)] if sources else "upload"),
            },
        )
        for i, chunk in enumerate(chunks)
    ]

    vs = QdrantVectorStore.from_existing_collection(
        embedding=_embeddings,
        collection_name=CUSTOM_COLLECTION,
        url=QDRANT_URL,
        api_key=QDRANT_API_KEY,
    )
    vs.add_documents(documents)
    return len(documents)


def get_program_vector_store() -> QdrantVectorStore:
    """Return a QdrantVectorStore pointed at the custom docs collection."""
    _ensure_custom_collection()
    return QdrantVectorStore.from_existing_collection(
        embedding=_embeddings,
        collection_name=CUSTOM_COLLECTION,
        url=QDRANT_URL,
        api_key=QDRANT_API_KEY,
    )
